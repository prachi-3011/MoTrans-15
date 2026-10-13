"""Export every final table to one Excel workbook (a convenience copy of the dataset).

Output: data/final/MoTrans15_all_records.xlsx (git-ignored: it contains licensed price data,
so it is shared privately like data/raw, never pushed to the public repository).

Sheets: README, panel_daily, policy_decisions, events, shocks, statements, controls_long,
hand_label_sample. Excel limits a cell to 32,767 characters, so long statement texts are cut
at 32,000 characters and flagged in `text_truncated`; the full texts are in statements.parquet.
"""

from __future__ import annotations

import datetime as dt
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
FINAL = ROOT / "data" / "final"
INTERIM = ROOT / "data" / "interim"
OUT = FINAL / "MoTrans15_all_records.xlsx"
MAX_TEXT = 32_000

SHEETS = [  # (sheet name, loader, description)
    ("panel_daily", lambda: pd.read_parquet(FINAL / "panel_daily.parquet"),
     "One row per trading day x equity index (38 series, 16 markets) with returns and controls."),
    ("policy_decisions", lambda: pd.read_csv(FINAL / "policy_decisions.csv"),
     "One row per central-bank decision (12 regimes, 2000-2026) with announcement windows."),
    ("events", lambda: pd.read_parquet(FINAL / "events.parquet"),
     "One row per decision x reacting index: matched sessions, event-window returns, shock."),
    ("shocks", lambda: pd.read_parquet(FINAL / "shocks.parquet"),
     "One row per decision: graded policy surprise (A high-frequency, B yield proxy, C none)."),
    ("statements", lambda: pd.read_parquet(FINAL / "statements.parquet"),
     "One row per English policy statement (Fed, ECB, BoJ, RBI); long texts truncated."),
    ("controls_long", lambda: pd.read_parquet(INTERIM / "controls_long.parquet"),
     "Every control observation (FX, policy rate, CPI, industrial production, 1y yield, VIX)."),
    ("hand_label_sample", lambda: pd.read_csv(FINAL / "hand_label_sample_with_text.csv"),
     "181 statements for hand labelling; label columns are intentionally empty."),
]

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3864")


def clean(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime()
    if value is pd.NaT:
        return None
    return value


def prepare(name: str, df: pd.DataFrame) -> pd.DataFrame:
    if "text" in df.columns:
        df = df.copy()
        df["text_truncated"] = df["text"].str.len() > MAX_TEXT
        df["text"] = df["text"].str.slice(0, MAX_TEXT)
    return df


def write_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    header = []
    for col in df.columns:
        cell = WriteOnlyCell(ws, value=str(col))
        cell.font, cell.fill = HEAD, FILL
        header.append(cell)
    ws.append(header)
    for i, col in enumerate(df.columns):
        width = 40 if col in ("text", "time_source", "date_source", "source_url") else 14
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = width
    for row in df.itertuples(index=False, name=None):
        ws.append([clean(v) for v in row])


def main() -> Path:
    wb = Workbook(write_only=True)
    loaded = [(name, prepare(name, loader()), desc) for name, loader, desc in SHEETS]

    readme = wb.create_sheet("README")
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 12
    readme.column_dimensions["C"].width = 110
    title = WriteOnlyCell(readme, value="MoTrans-15 dataset - all records")
    title.font = Font(name="Arial", size=14, bold=True)
    readme.append([title])
    readme.append([WriteOnlyCell(readme, value=f"Exported {dt.date.today().isoformat()} by src/export_excel.py "
                                               "from data/final/ and data/interim/.")])
    readme.append([WriteOnlyCell(readme, value="Contains licensed index data (Yahoo, NSE, CSI, MOEX, B3): "
                                               "for the project team only - do not publish or redistribute.")])
    readme.append([])
    hdr = []
    for text in ("Sheet", "Rows", "What one row is"):
        c = WriteOnlyCell(readme, value=text)
        c.font, c.fill = HEAD, FILL
        hdr.append(c)
    readme.append(hdr)
    for name, df, desc in loaded:
        cells = [WriteOnlyCell(readme, value=v) for v in (name, len(df), desc)]
        for c in cells:
            c.font, c.alignment = FONT, Alignment(wrap_text=True, vertical="top")
        readme.append(cells)
    readme.append([])
    readme.append([WriteOnlyCell(readme, value="Column definitions: docs/data_dictionary.md. "
                                               "Known gaps: docs/issues.md. Choices: DECISIONS.md.")])

    for name, df, _ in loaded:
        write_sheet(wb, name, df)
        print(f"{name}: {len(df):,} rows x {df.shape[1]} columns")
    wb.save(OUT)
    print(f"Saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB)")
    return OUT


if __name__ == "__main__":
    main()
